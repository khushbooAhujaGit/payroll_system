"""Dump structure of both sheets in the Leave Tracker xlsm."""
import openpyxl, re

FILE = r'C:\Users\Shivam\Downloads\01 Leave Tracker.xlsm'

wb = openpyxl.load_workbook(FILE, read_only=True, keep_vba=True, data_only=True)
print(f"Sheets: {wb.sheetnames}\n")

for sh_name in wb.sheetnames:
    ws = wb[sh_name]
    print(f"\n{'='*80}")
    print(f"SHEET: {sh_name}  (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"{'='*80}")
    rows_shown = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        # skip fully-empty rows
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        print(f"  row{i+1:03d}: {[str(v)[:30] if v is not None else None for v in row]}")
        rows_shown += 1
        if rows_shown >= 60:
            print("  ... (truncated)")
            break

wb.close()
print("\nDone.")
