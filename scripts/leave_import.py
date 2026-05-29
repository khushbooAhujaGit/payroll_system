#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Leave Tracker Import  --  HRMS Payroll
======================================
Reads the "Mark Leave" sheet from the Leave Tracker xlsm and inserts records
into leave_tracker_records.

Usage
-----
  python leave_import.py <01 Leave Tracker.xlsm>  [--dry-run]

Arguments
---------
  <file.xlsm>   Path to the Leave Tracker Excel file (xlsm or xlsx)
  --dry-run     Parse and validate without writing to the database.
                Prints summary and opens error report if issues found.

Rules
-----
  * Reads the "Mark Leave" sheet (tab name must match exactly).
  * Row 1 is treated as headers — skipped.
  * Columns: ID | EMP ID | Name | Date | Leave Type | Team LeaderName | Reason Type | Reason
  * EMP ID is matched against employees.employee_id in the database.
    Rows with unmatched EMP IDs are imported with employee_id = NULL
    (no error — just a warning row in the report).
  * Missing Date or Leave Type = hard error (row goes to error report).
  * ON DUPLICATE KEY UPDATE ensures re-running is idempotent.
  * Error report is an xlsx file in the script directory and is opened
    automatically when any errors are found.
"""

from __future__ import annotations  # allows X | Y union types on Python 3.9

import os
import re
import sys
import subprocess
import datetime
from pathlib import Path

import openpyxl
import xlsxwriter
import mysql.connector

# Force UTF-8 output on Windows (avoids UnicodeEncodeError for box-drawing chars)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

# ── Bootstrap ─────────────────────────────────────────────────────────────────

SCRIPT_DIR   = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, '..'))

_env_path = os.path.join(PROJECT_ROOT, '.env')
if os.path.isfile(_env_path):
    with open(_env_path, encoding='utf-8') as _fh:
        for _line in _fh:
            _line = _line.strip()
            if _line and not _line.startswith('#') and '=' in _line:
                _k, _, _v = _line.partition('=')
                os.environ.setdefault(_k.strip(), _v.strip().strip('"').strip("'"))

DB_CONFIG = {
    'host':       os.getenv('DB_HOST',     '127.0.0.1'),
    'port':       int(os.getenv('DB_PORT', '3306')),
    'user':       os.getenv('DB_USERNAME', 'root'),
    'password':   os.getenv('DB_PASSWORD', ''),
    'database':   os.getenv('DB_DATABASE', 'hrms'),
    'charset':    'utf8mb4',
    'autocommit': False,
}

SHEET_NAME  = 'Mark Leave'
ERROR_FILE  = os.path.join(SCRIPT_DIR, 'leave_import_errors.xlsx')

# ── Column header → field name map ────────────────────────────────────────────

HEADER_MAP = {
    'id':               'leave_ref_id',
    'emp id':           'sheet_emp_id',
    'name':             'employee_name',
    'date':             'leave_date',
    'leave type':       'leave_type',
    'team leadername':  'team_leader_name',
    'reason type':      'reason_type',
    'reason':           'reason',
}

# ── Helpers ────────────────────────────────────────────────────────────────────

def cell_str(value) -> str:
    """Convert a cell value to a stripped string; return '' for None/empty."""
    if value is None:
        return ''
    return str(value).strip()


def cell_date(value) -> datetime.date | None:
    """
    Convert a cell value to a date.
    Handles: datetime, date, 'YYYY-MM-DD', 'DD-MM-YYYY', 'MM/DD/YYYY', etc.
    Returns None on failure.
    """
    if value is None or value == '':
        return None

    if isinstance(value, (datetime.datetime, datetime.date)):
        if isinstance(value, datetime.datetime):
            return value.date()
        return value

    s = str(value).strip()
    if not s:
        return None

    # Try common date formats
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y', '%d/%m/%Y',
                '%d-%b-%Y', '%Y/%m/%d', '%d.%m.%Y'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass

    # ISO-ish with time part
    try:
        return datetime.date.fromisoformat(s[:10])
    except (ValueError, IndexError):
        pass

    return None


def open_file(path: str) -> None:
    """Open a file with the OS default application."""
    try:
        if sys.platform == 'win32':
            os.startfile(path)
        elif sys.platform == 'darwin':
            subprocess.call(['open', path])
        else:
            subprocess.call(['xdg-open', path])
    except Exception:
        pass


# ── Error report ───────────────────────────────────────────────────────────────

def write_error_report(errors: list[dict]) -> None:
    """
    Write an xlsx error report.  Each item in `errors` has:
        row_num, leave_ref_id, sheet_emp_id, employee_name, leave_date_raw,
        leave_type, team_leader_name, reason_type, reason, error_msg
    """
    wb  = xlsxwriter.Workbook(ERROR_FILE)
    ws  = wb.add_worksheet('Errors')

    hdr_fmt  = wb.add_format({'bold': True, 'bg_color': '#C0392B', 'font_color': '#FFFFFF',
                               'border': 1, 'text_wrap': True})
    err_fmt  = wb.add_format({'bg_color': '#FDE8E8', 'border': 1, 'text_wrap': True, 'valign': 'top'})
    warn_fmt = wb.add_format({'bg_color': '#FFF8E1', 'border': 1, 'text_wrap': True, 'valign': 'top'})

    headers = ['Row #', 'ID', 'EMP ID', 'Name', 'Date', 'Leave Type',
               'Team Leader', 'Reason Type', 'Reason', 'Error / Warning']
    col_widths = [8, 18, 14, 25, 14, 12, 22, 18, 30, 50]

    for c, (h, w) in enumerate(zip(headers, col_widths)):
        ws.write(0, c, h, hdr_fmt)
        ws.set_column(c, c, w)

    for r, e in enumerate(errors, start=1):
        fmt = warn_fmt if e.get('is_warning') else err_fmt
        ws.write(r, 0, e.get('row_num', ''),           fmt)
        ws.write(r, 1, e.get('leave_ref_id', ''),      fmt)
        ws.write(r, 2, e.get('sheet_emp_id', ''),      fmt)
        ws.write(r, 3, e.get('employee_name', ''),     fmt)
        ws.write(r, 4, e.get('leave_date_raw', ''),    fmt)
        ws.write(r, 5, e.get('leave_type', ''),        fmt)
        ws.write(r, 6, e.get('team_leader_name', ''),  fmt)
        ws.write(r, 7, e.get('reason_type', ''),       fmt)
        ws.write(r, 8, e.get('reason', ''),            fmt)
        ws.write(r, 9, e.get('error_msg', ''),         fmt)

    wb.close()
    print(f"\n  Error/warning report saved → {ERROR_FILE}")
    open_file(ERROR_FILE)


# ── Main parsing ───────────────────────────────────────────────────────────────

def parse_sheet(filepath: str) -> tuple[list[dict], list[dict]]:
    """
    Load the xlsm/xlsx file, read the 'Mark Leave' sheet, and return:
        (records, error_rows)

    records    — list of dicts ready for DB insert (all fields populated)
    error_rows — list of rows with hard errors (missing date or leave type)
    """
    print(f"  Loading workbook: {filepath}")
    wb = openpyxl.load_workbook(filepath, read_only=True, keep_vba=True, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        print(f"  [ERROR] Sheet '{SHEET_NAME}' not found.")
        print(f"  Available sheets: {wb.sheetnames}")
        wb.close()
        sys.exit(1)

    ws = wb[SHEET_NAME]
    print(f"  Sheet '{SHEET_NAME}': max_row={ws.max_row}, max_col={ws.max_column}")

    # ── Map header row ──────────────────────────────────────────────────────
    col_map = {}   # col_index (0-based) → field name
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        print("  [ERROR] Sheet is empty.")
        wb.close()
        sys.exit(1)

    for i, h in enumerate(header_row):
        key = str(h).strip().lower() if h is not None else ''
        field = HEADER_MAP.get(key)
        if field:
            col_map[i] = field

    print(f"  Mapped columns: {col_map}")

    required_fields = {'leave_ref_id', 'leave_date', 'leave_type'}
    found_fields    = set(col_map.values())
    missing = required_fields - found_fields
    if missing:
        print(f"  [ERROR] Required columns not found in sheet headers: {missing}")
        wb.close()
        sys.exit(1)

    # ── Parse data rows ─────────────────────────────────────────────────────
    records    = []
    error_rows = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # Skip entirely blank rows
        if all(v is None or str(v).strip() == '' for v in row):
            continue

        # Extract fields
        r: dict = {field: '' for field in HEADER_MAP.values()}
        for col_i, field in col_map.items():
            if col_i < len(row):
                val = row[col_i]
                r[field] = val  # keep original value for date fields

        leave_ref_id      = cell_str(r['leave_ref_id'])
        sheet_emp_id      = cell_str(r['sheet_emp_id'])
        employee_name     = cell_str(r['employee_name'])
        leave_date_raw    = r['leave_date']
        leave_type        = cell_str(r['leave_type']).upper()
        team_leader_name  = cell_str(r['team_leader_name']) or None
        reason_type       = cell_str(r['reason_type']) or None
        reason            = cell_str(r['reason']) or None

        # Parse date
        leave_date = cell_date(leave_date_raw)

        # Hard error checks
        row_errors = []
        if not leave_ref_id:
            row_errors.append('ID (leave reference) is required')
        if leave_date is None:
            raw_str = cell_str(leave_date_raw) or repr(leave_date_raw)
            row_errors.append(f'Date is required and must be a valid date (got: {raw_str!r})')
        if not leave_type:
            row_errors.append('Leave Type is required')

        if row_errors:
            error_rows.append({
                'row_num':          row_idx,
                'leave_ref_id':     leave_ref_id,
                'sheet_emp_id':     sheet_emp_id,
                'employee_name':    employee_name,
                'leave_date_raw':   cell_str(leave_date_raw),
                'leave_type':       leave_type,
                'team_leader_name': team_leader_name or '',
                'reason_type':      reason_type or '',
                'reason':           reason or '',
                'error_msg':        '; '.join(row_errors),
                'is_warning':       False,
            })
            continue

        records.append({
            'leave_ref_id':     leave_ref_id,
            'sheet_emp_id':     sheet_emp_id or None,
            'employee_name':    employee_name or None,
            'leave_date':       leave_date,
            'leave_type':       leave_type,
            'team_leader_name': team_leader_name,
            'reason_type':      reason_type,
            'reason':           reason,
        })

    wb.close()
    return records, error_rows


# ── Database operations ────────────────────────────────────────────────────────

def resolve_employee_ids(cursor, records: list[dict]) -> dict[str, int | None]:
    """
    Batch-resolve sheet EMP IDs to employees.id.
    Returns dict: sheet_emp_id → employees.id (or None)
    """
    emp_ids = list({r['sheet_emp_id'] for r in records if r['sheet_emp_id']})
    if not emp_ids:
        return {}

    placeholders = ','.join(['%s'] * len(emp_ids))
    cursor.execute(
        f"SELECT employee_id, id FROM employees WHERE employee_id IN ({placeholders})",
        emp_ids
    )
    return {row[0]: row[1] for row in cursor.fetchall()}


def insert_records(cursor, records: list[dict], emp_map: dict,
                   source_file: str, uploaded_by: int | None) -> tuple[int, list[dict]]:
    """
    Upsert records into leave_tracker_records.
    Returns (inserted_count, warning_rows).
    warning_rows are rows where employee_id could not be resolved.
    """
    SQL = """
        INSERT INTO leave_tracker_records
            (leave_ref_id, employee_id, sheet_emp_id, employee_name,
             leave_date, leave_type, team_leader_name, reason_type, reason,
             source_file, uploaded_by, created_at, updated_at)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW(), NOW())
        ON DUPLICATE KEY UPDATE
            employee_id      = VALUES(employee_id),
            sheet_emp_id     = VALUES(sheet_emp_id),
            employee_name    = VALUES(employee_name),
            leave_date       = VALUES(leave_date),
            leave_type       = VALUES(leave_type),
            team_leader_name = VALUES(team_leader_name),
            reason_type      = VALUES(reason_type),
            reason           = VALUES(reason),
            source_file      = VALUES(source_file),
            uploaded_by      = VALUES(uploaded_by),
            updated_at       = NOW()
    """

    inserted  = 0
    warnings  = []

    for rec in records:
        emp_id = emp_map.get(rec['sheet_emp_id']) if rec['sheet_emp_id'] else None

        if rec['sheet_emp_id'] and emp_id is None:
            warnings.append({
                'row_num':          '—',
                'leave_ref_id':     rec['leave_ref_id'],
                'sheet_emp_id':     rec['sheet_emp_id'],
                'employee_name':    rec['employee_name'] or '',
                'leave_date_raw':   str(rec['leave_date']),
                'leave_type':       rec['leave_type'],
                'team_leader_name': rec['team_leader_name'] or '',
                'reason_type':      rec['reason_type'] or '',
                'reason':           rec['reason'] or '',
                'error_msg':        f'EMP ID "{rec["sheet_emp_id"]}" not found in employees — imported with employee_id = NULL',
                'is_warning':       True,
            })

        cursor.execute(SQL, (
            rec['leave_ref_id'],
            emp_id,
            rec['sheet_emp_id'],
            rec['employee_name'],
            rec['leave_date'].strftime('%Y-%m-%d'),
            rec['leave_type'],
            rec['team_leader_name'],
            rec['reason_type'],
            rec['reason'],
            source_file,
            uploaded_by,
        ))
        inserted += 1

    return inserted, warnings


# ── Entry point ────────────────────────────────────────────────────────────────

def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    dry_run = '--dry-run' in sys.argv

    if len(args) != 1:
        print("Usage: python leave_import.py <Leave Tracker.xlsm> [--dry-run]")
        sys.exit(1)

    filepath = os.path.abspath(args[0])
    if not os.path.isfile(filepath):
        print(f"[ERROR] File not found: {filepath}")
        sys.exit(1)

    source_file  = os.path.basename(filepath)
    uploaded_by  = None   # no web user in CLI context

    print("=" * 70)
    print("HRMS — Leave Tracker Import")
    print("=" * 70)
    print(f"  File     : {filepath}")
    print(f"  Dry run  : {'YES — no DB writes' if dry_run else 'NO'}")
    print()

    # ── Parse ──────────────────────────────────────────────────────────────
    records, error_rows = parse_sheet(filepath)

    print(f"\n  Parsed    : {len(records) + len(error_rows):,} data rows")
    print(f"  Valid     : {len(records):,}")
    print(f"  Errors    : {len(error_rows):,}")

    if error_rows:
        print(f"\n  [FAIL] {len(error_rows)} row(s) have validation errors:")
        for e in error_rows[:10]:
            print(f"    Row {e['row_num']:>5}: {e['error_msg']}")
        if len(error_rows) > 10:
            print(f"    ... and {len(error_rows) - 10} more")
        write_error_report(error_rows)
        print("\n  Import aborted — fix the errors and re-run.")
        sys.exit(1)

    if not records:
        print("\n  No valid data rows found. Nothing to import.")
        sys.exit(0)

    if dry_run:
        # Show a sample
        print("\n  Sample records (first 5):")
        for r in records[:5]:
            print(f"    {r['leave_ref_id']:<18} {r['sheet_emp_id']:<12} "
                  f"{str(r['leave_date']):<12} {r['leave_type']:<6} {r['employee_name']}")
        print("\n  [DRY RUN] No database writes performed.")
        sys.exit(0)

    # ── Database ────────────────────────────────────────────────────────────
    print("\n  Connecting to database …")
    conn   = mysql.connector.connect(**DB_CONFIG)
    cursor = conn.cursor()

    try:
        print("  Resolving employee IDs …")
        emp_map = resolve_employee_ids(cursor, records)

        unmatched_ids = {r['sheet_emp_id'] for r in records
                         if r['sheet_emp_id'] and emp_map.get(r['sheet_emp_id']) is None}
        if unmatched_ids:
            print(f"  [WARN] {len(unmatched_ids)} EMP ID(s) not found in employees table "
                  f"— those rows import with employee_id = NULL:")
            for uid in sorted(unmatched_ids)[:10]:
                print(f"    {uid}")
            if len(unmatched_ids) > 10:
                print(f"    … and {len(unmatched_ids) - 10} more")

        print(f"  Inserting / updating {len(records):,} records …")
        inserted, warnings = insert_records(cursor, records, emp_map, source_file, uploaded_by)
        conn.commit()

        print(f"\n  [OK] {inserted:,} leave record(s) imported successfully.")

        if warnings:
            print(f"  Writing warning report ({len(warnings)} unmatched EMP IDs) …")
            write_error_report(warnings)
        else:
            print("  No warnings — all EMP IDs resolved.")

    except Exception as e:
        conn.rollback()
        print(f"\n  [ERROR] Database error: {e}")
        raise
    finally:
        cursor.close()
        conn.close()

    print("=" * 70)
    print("DONE")
    print("=" * 70)


if __name__ == '__main__':
    main()
